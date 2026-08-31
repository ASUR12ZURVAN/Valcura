import openpyxl
import os
import datetime as dt
from typing import Dict, List, Tuple, Any
from decimal import Decimal
from .models import (
    ClinicProfile, TeamMaster, PatientContact, Opportunity, 
    InteractionLog, TreatmentConfig, SystemConfig, MonthlyMetrics
)


class ExcelImporterService:
    """
    Service to import data from Valcura Master Revenue Intelligence Database Excel file
    into Django models.
    """

    def __init__(self, excel_file_path: str = None):
        if excel_file_path is None:
            # Default path relative to project root
            excel_file_path = os.path.join(
                os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
                'Valcura - Master Revenue Intelligence Database.xlsx'
            )
        self.excel_file_path = excel_file_path
        self.workbook = None

    def load_workbook(self):
        """Load the Excel workbook."""
        if not os.path.exists(self.excel_file_path):
            raise FileNotFoundError(f"Excel file not found: {self.excel_file_path}")
        self.workbook = openpyxl.load_workbook(self.excel_file_path)

    def import_all_sheets(self) -> Dict[str, Any]:
        """Import all sheets from the Excel file."""
        self.load_workbook()
        
        results = {
            'success': True,
            'imported': {},
            'errors': []
        }

        try:
            # Import in dependency order
            results['imported']['Clinic_Profile'] = self.import_clinic_profile()
            results['imported']['Team_Master'] = self.import_team_master()
            results['imported']['Patient_Contact'] = self.import_patient_contact()
            results['imported']['Opportunities'] = self.import_opportunities()
            results['imported']['Interaction_Log'] = self.import_interaction_log()
            results['imported']['Treatment_Config'] = self.import_treatment_config()
            results['imported']['System_Config'] = self.import_system_config()
            results['imported']['Monthly_Metrics'] = self.import_monthly_metrics()

        except Exception as e:
            results['success'] = False
            results['errors'].append(str(e))

        return results

    def import_clinic_profile(self) -> Dict[str, int]:
        """Import Clinic_Profile sheet."""
        sheet = self.workbook['Clinic_Profile']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Skip empty rows
                continue

            try:
                clinic_data = {
                    'clinic_id': str(row[0]),
                    'clinic_name': row[1],
                    'city': row[2],
                    'locality': row[3],
                    'clinic_type': row[4],
                    'chair_count': int(row[5]) if row[5] else 0,
                    'primary_specialty': row[6],
                    'monthly_revenue_baseline_inr': Decimal(str(row[7])) if row[7] else Decimal('0'),
                    'reception_team_size': int(row[8]) if row[8] else 0,
                    'onboarding_date': self._parse_date(row[9]),
                    'crm_status': row[10],
                    'crm_name': row[11] if row[11] else '',
                    'crm_integration_mode': row[12] if row[12] else '',
                    'valcura_plan': row[13],
                    'active_status': bool(row[14]) if row[14] is not None else True,
                }

                clinic, created = ClinicProfile.objects.update_or_create(
                    clinic_id=clinic_data['clinic_id'],
                    defaults=clinic_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing clinic profile row {row[0]}: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def import_team_master(self) -> Dict[str, int]:
        """Import Team_Master sheet."""
        sheet = self.workbook['Team_Master']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue

            try:
                clinic = ClinicProfile.objects.filter(clinic_id=str(row[1])).first()
                if not clinic:
                    print(f"Clinic {row[1]} not found for staff {row[0]}")
                    continue

                team_data = {
                    'staff_id': str(row[0]),
                    'clinic': clinic,
                    'staff_name': row[2],
                    'role': row[3],
                    'whatsapp_number': row[4],
                    'active_status': bool(row[5]) if row[5] is not None else True,
                    'start_date': self._parse_date(row[6]),
                }

                staff, created = TeamMaster.objects.update_or_create(
                    staff_id=team_data['staff_id'],
                    defaults=team_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing team master row {row[0]}: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def import_patient_contact(self) -> Dict[str, int]:
        """Import Patient_Contact sheet."""
        sheet = self.workbook['Patient_Contact']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue

            try:
                clinic = ClinicProfile.objects.filter(clinic_id=str(row[1])).first()
                if not clinic:
                    print(f"Clinic {row[1]} not found for patient {row[0]}")
                    continue

                patient_data = {
                    'patient_id': str(row[0]),
                    'clinic': clinic,
                    'external_crm_id': row[2] if row[2] else '',
                    'patient_name': row[3],
                    'phone_number': row[4],
                    'phone_hash_analytics': row[5] if row[5] else '',
                    'new_existing_patient': row[6],
                    'preferred_language': row[7] if row[7] else 'English',
                    'communication_allowed': bool(row[8]) if row[8] is not None else True,
                    'data_source': row[9],
                    'created_at': self._parse_datetime(row[10]),
                    'updated_at': self._parse_datetime(row[11]),
                }

                patient, created = PatientContact.objects.update_or_create(
                    patient_id=patient_data['patient_id'],
                    defaults=patient_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing patient contact row {row[0]}: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def import_opportunities(self) -> Dict[str, int]:
        """Import Opportunities sheet."""
        sheet = self.workbook['Opportunities']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue

            try:
                patient = PatientContact.objects.filter(patient_id=str(row[1])).first()
                clinic = ClinicProfile.objects.filter(clinic_id=str(row[2])).first()
                doctor = TeamMaster.objects.filter(staff_id=str(row[10])).first() if row[10] else None

                if not patient or not clinic:
                    print(f"Patient or clinic not found for opportunity {row[0]}")
                    continue

                opportunity_data = {
                    'opportunity_id': str(row[0]),
                    'patient': patient,
                    'clinic': clinic,
                    'external_crm_opportunity_id': row[3] if row[3] else '',
                    'inquiry_date_time': self._parse_datetime(row[4]),
                    'inquiry_type': row[5],
                    'lead_source': row[6],
                    'treatment_concern': row[7],
                    'appointment_date_time': self._parse_datetime(row[8]),
                    'consultation_date_time': self._parse_datetime(row[9]),
                    'doctor': doctor,
                    'treatment_advised': row[11] if row[11] else '',
                    'why_treatment_advised': row[12] if row[12] else '',
                    'ats_value_inr': Decimal(str(row[13])) if row[13] else None,
                    'ats_category': row[14] if row[14] else '',
                    'urgency_level': row[15] if row[15] else '',
                    'initial_objection': row[16] if row[16] else '',
                    'current_objection': row[17] if row[17] else '',
                    'opportunity_status': row[18],
                    'last_interaction_date_time': self._parse_datetime(row[19]),
                    'next_followup_date_time': self._parse_datetime(row[20]),
                    'treatment_start_date': self._parse_date(row[21]),
                    'treatment_completion_date': self._parse_date(row[22]),
                    'realized_revenue_inr': Decimal(str(row[23])) if row[23] else Decimal('0'),
                    'lost_reason': row[24] if row[24] else '',
                    'nps_score': int(row[25]) if row[25] else None,
                    'conversion_probability_pct': Decimal(str(row[26])) if row[26] else None,
                    'weighted_pipeline_value_inr': Decimal(str(row[27])) if row[27] else Decimal('0'),
                    'revenue_at_risk_inr': Decimal(str(row[28])) if row[28] else Decimal('0'),
                    'conversion_days': int(row[29]) if row[29] else None,
                }

                opportunity, created = Opportunity.objects.update_or_create(
                    opportunity_id=opportunity_data['opportunity_id'],
                    defaults=opportunity_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing opportunity row {row[0]}: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def import_interaction_log(self) -> Dict[str, int]:
        """Import Interaction_Log sheet."""
        sheet = self.workbook['Interaction_Log']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue

            try:
                opportunity = Opportunity.objects.filter(opportunity_id=str(row[1])).first()
                patient = PatientContact.objects.filter(patient_id=str(row[2])).first()
                clinic = ClinicProfile.objects.filter(clinic_id=str(row[3])).first()
                staff = TeamMaster.objects.filter(staff_id=str(row[4])).first() if row[4] else None

                if not patient or not clinic:
                    print(f"Patient or clinic not found for interaction {row[0]}")
                    continue

                interaction_data = {
                    'interaction_id': str(row[0]),
                    'opportunity': opportunity,
                    'patient': patient,
                    'clinic': clinic,
                    'staff': staff,
                    'event_date_time': self._parse_datetime(row[5]),
                    'channel': row[6],
                    'direction': row[7],
                    'interaction_type': row[8],
                    'sequence_name': row[9] if row[9] else '',
                    'day_in_sequence': int(row[10]) if row[10] else None,
                    'message_variant_id': row[11] if row[11] else '',
                    'call_objective': row[12] if row[12] else '',
                    'outcome': row[13] if row[13] else '',
                    'response_status': row[14] if row[14] else '',
                    'previous_objection': row[15] if row[15] else '',
                    'observed_objection': row[16] if row[16] else '',
                    'next_followup_date_time': self._parse_datetime(row[17]),
                    'compliance_flag': bool(row[18]) if row[18] is not None else None,
                    'delivery_status': row[19] if row[19] else '',
                    'ai_summary': row[20] if row[20] else '',
                    'conversion_fingerprint_flag': bool(row[21]) if row[21] is not None else False,
                    'notes': row[22] if row[22] else '',
                }

                interaction, created = InteractionLog.objects.update_or_create(
                    interaction_id=interaction_data['interaction_id'],
                    defaults=interaction_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing interaction log row {row[0]}: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def import_treatment_config(self) -> Dict[str, int]:
        """Import Treatment_Config sheet."""
        sheet = self.workbook['Treatment_Config']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[1]:  # Treatment is required
                continue

            try:
                clinic = ClinicProfile.objects.filter(clinic_id=str(row[0])).first()
                if not clinic:
                    print(f"Clinic {row[0]} not found for treatment config")
                    continue

                treatment_data = {
                    'clinic': clinic,
                    'treatment': row[1],
                    'default_ats_value_inr': Decimal(str(row[2])) if row[2] else Decimal('0'),
                    'ats_category': row[3] if row[3] else '',
                    'expected_sittings': int(row[4]) if row[4] else 1,
                    'followup_track': row[5] if row[5] else '',
                    'default_urgency': row[6] if row[6] else '',
                    'active_status': bool(row[7]) if row[7] is not None else True,
                    'created_at': self._parse_datetime(row[8]),
                    'updated_at': self._parse_datetime(row[9]),
                }

                treatment_config, created = TreatmentConfig.objects.update_or_create(
                    clinic=treatment_data['clinic'],
                    treatment=treatment_data['treatment'],
                    defaults=treatment_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing treatment config row: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def import_system_config(self) -> Dict[str, int]:
        """Import System_Config sheet."""
        sheet = self.workbook['System_Config']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[1]:  # Code is required
                continue

            try:
                system_data = {
                    'config_type': row[0],
                    'code': row[1],
                    'display_value': row[2],
                    'treatment': row[3] if row[3] else '',
                    'objection': row[4] if row[4] else '',
                    'ats_category': row[5] if row[5] else '',
                    'language': row[6] if row[6] else 'English',
                    'sequence_name': row[7],
                    'day_in_sequence': int(row[8]) if row[8] else 1,
                    'meta_template_name': row[9],
                    'active_status': bool(row[10]) if row[10] is not None else True,
                    'version': row[11] if row[11] else 'v1',
                }

                config, created = SystemConfig.objects.update_or_create(
                    code=system_data['code'],
                    version=system_data['version'],
                    defaults=system_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing system config row {row[1]}: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def import_monthly_metrics(self) -> Dict[str, int]:
        """Import Monthly_Metrics sheet."""
        sheet = self.workbook['Monthly_Metrics']
        imported = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Month is required
                continue

            try:
                clinic = ClinicProfile.objects.filter(clinic_id=str(row[1])).first()
                if not clinic:
                    print(f"Clinic {row[1]} not found for monthly metrics")
                    continue

                month_date = self._parse_date(row[0])
                if not month_date:
                    continue

                metrics_data = {
                    'month': month_date,
                    'clinic': clinic,
                    'confirmed_revenue_inr': Decimal(str(row[2])) if row[2] else Decimal('0'),
                    'inquiry_count': int(row[3]) if row[3] else 0,
                    'consultation_count': int(row[4]) if row[4] else 0,
                    'treatment_started_count': int(row[5]) if row[5] else 0,
                    'inquiry_to_consult_pct': Decimal(str(row[6])) if row[6] else Decimal('0'),
                    'consult_to_treatment_pct': Decimal(str(row[7])) if row[7] else Decimal('0'),
                    'missed_call_recovery_pct': Decimal(str(row[8])) if row[8] else Decimal('0'),
                    'followup_compliance_pct': Decimal(str(row[9])) if row[9] else Decimal('0'),
                    'pipeline_value_inr': Decimal(str(row[10])) if row[10] else Decimal('0'),
                    'revenue_at_risk_inr': Decimal(str(row[11])) if row[11] else Decimal('0'),
                    'recovered_revenue_inr': Decimal(str(row[12])) if row[12] else Decimal('0'),
                    'avg_nps': Decimal(str(row[13])) if row[13] else None,
                    'data_quality_pct': Decimal(str(row[14])) if row[14] else Decimal('0'),
                    'forecast_30d_inr': Decimal(str(row[15])) if row[15] else Decimal('0'),
                    'forecast_60d_inr': Decimal(str(row[16])) if row[16] else Decimal('0'),
                    'forecast_90d_inr': Decimal(str(row[17])) if row[17] else Decimal('0'),
                }

                metrics, created = MonthlyMetrics.objects.update_or_create(
                    month=metrics_data['month'],
                    clinic=metrics_data['clinic'],
                    defaults=metrics_data
                )

                if created:
                    imported += 1
                else:
                    updated += 1

            except Exception as e:
                print(f"Error importing monthly metrics row: {str(e)}")

        return {'imported': imported, 'updated': updated}

    def _parse_date(self, date_value) -> dt.date:
        """Parse date from Excel cell value."""
        if date_value is None:
            return None
        if isinstance(date_value, dt.datetime):
            return date_value.date()
        if isinstance(date_value, str):
            try:
                return dt.datetime.strptime(date_value, '%Y-%m-%d').date()
            except ValueError:
                try:
                    return dt.datetime.strptime(date_value, '%d/%m/%Y').date()
                except ValueError:
                    return None
        if isinstance(date_value, (int, float)):
            return dt.datetime.fromtimestamp(dt.datetime(1899, 12, 30).timestamp() + date_value * 86400).date()
        return None

    def _parse_datetime(self, datetime_value) -> dt.datetime:
        """Parse datetime from Excel cell value."""
        if datetime_value is None:
            return None
        if isinstance(datetime_value, dt.datetime):
            return datetime_value
        if isinstance(datetime_value, str):
            try:
                return dt.datetime.strptime(datetime_value, '%Y-%m-%d %H:%M:%S')
            except ValueError:
                try:
                    return dt.datetime.strptime(datetime_value, '%Y-%m-%d %H:%M')
                except ValueError:
                    try:
                        return dt.datetime.strptime(datetime_value, '%d/%m/%Y %H:%M:%S')
                    except ValueError:
                        return None
        if isinstance(datetime_value, (int, float)):
            return dt.datetime.fromtimestamp(dt.datetime(1899, 12, 30).timestamp() + datetime_value * 86400)
        return None
